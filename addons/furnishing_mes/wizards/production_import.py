# -*- coding: utf-8 -*-
"""DAY WISE OUTPUT importer.

Brings the customer's existing production spreadsheet into the system. This is
both the migration path for their history — which is what makes trend analysis
useful at go-live rather than six months later — and the catch-up route for a
day the terminal was not used.

The column mapping is **configuration, not code**. The real layout of the
customer's workbook is still open (question Q8), so the wizard reads the header
row from the uploaded file and lets the user say which column means what. When
the real file arrives it is a mapping choice, not a rework.

Every import is a batch that can be reversed as a unit.
"""

import base64
import io
import logging
from datetime import datetime

from odoo import _, api, fields, models
from odoo.exceptions import UserError

_logger = logging.getLogger(__name__)

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - openpyxl ships with Odoo
    load_workbook = None


class FmesProductionImport(models.TransientModel):
    _name = 'fmes.production.import'
    _description = 'Import Daily Production Output'

    file_data = fields.Binary(string='Spreadsheet', required=True)
    file_name = fields.Char(string='File Name')
    sheet_name = fields.Char(
        string='Sheet',
        help="Leave empty to read the first sheet in the workbook.")
    header_row = fields.Integer(
        string='Header Row', default=1, required=True,
        help="Row number holding the column titles. Data is read from the "
             "row after it.")

    state = fields.Selection(
        [('upload', 'Upload'), ('map', 'Map Columns'), ('preview', 'Preview')],
        default='upload', required=True)

    detected_columns = fields.Char(readonly=True)

    # ------------------------------------------------------- column mapping
    col_date = fields.Char(string='Date Column', help="Required.")
    col_shift = fields.Char(string='Shift Column')
    col_machine = fields.Char(string='Machine Column', help="Required.")
    col_product = fields.Char(string='Product Column', help="Required.")
    col_planned = fields.Char(string='Target Column')
    col_actual = fields.Char(string='Produced Column', help="Required.")
    col_rejected = fields.Char(string='Rejected Column')
    col_run_hours = fields.Char(string='Run Hours Column')
    col_downtime = fields.Char(string='Downtime Hours Column')
    col_manpower = fields.Char(string='Manpower Column')
    col_remarks = fields.Char(string='Remarks Column')

    default_shift_id = fields.Many2one(
        'fmes.shift', string='Default Shift',
        help="Used for rows with no shift column or an unrecognised value.")
    match_machine_by = fields.Selection(
        [('code', 'Machine Code'), ('name', 'Machine Name')],
        default='code', required=True)
    match_product_by = fields.Selection(
        [('default_code', 'Internal Reference'), ('name', 'Product Name')],
        default='default_code', required=True)
    create_missing_entries = fields.Boolean(
        string='Import Rows Without a Plan', default=True,
        help="Import output for slots that were never planned. Usually wanted "
             "when migrating history, since the old sheets predate the plans.")

    # -------------------------------------------------------------- preview
    preview_html = fields.Html(readonly=True, sanitize=False)
    row_count = fields.Integer(readonly=True)
    valid_count = fields.Integer(readonly=True)
    error_count = fields.Integer(readonly=True)
    company_id = fields.Many2one(
        'res.company', required=True, default=lambda self: self.env.company)

    # ==================================================================
    # Reading the workbook
    # ==================================================================
    def _load_rows(self):
        """Return (headers, rows) from the uploaded workbook."""
        self.ensure_one()
        if load_workbook is None:
            raise UserError(_("openpyxl is not available on this server."))
        if not self.file_data:
            raise UserError(_("Upload a file first."))

        try:
            content = base64.b64decode(self.file_data)
            workbook = load_workbook(
                io.BytesIO(content), data_only=True, read_only=True)
        except Exception as exc:
            raise UserError(_(
                "That file could not be read as an Excel workbook (.xlsx). "
                "If it is an old .xls file, open it in Excel and save it as "
                ".xlsx first.\n\n%s", exc))

        if self.sheet_name and self.sheet_name in workbook.sheetnames:
            sheet = workbook[self.sheet_name]
        else:
            sheet = workbook[workbook.sheetnames[0]]

        rows = list(sheet.iter_rows(values_only=True))
        header_index = max(self.header_row - 1, 0)
        if header_index >= len(rows):
            raise UserError(_(
                "The sheet has %(rows)s rows, so row %(header)s cannot be the "
                "header.", rows=len(rows), header=self.header_row))

        headers = [
            (str(cell).strip() if cell is not None else '')
            for cell in rows[header_index]
        ]
        data_rows = rows[header_index + 1:]
        return headers, data_rows

    def action_read_headers(self):
        """Step 1: read the file and show what columns it has."""
        self.ensure_one()
        headers, rows = self._load_rows()
        present = [h for h in headers if h]
        if not present:
            raise UserError(_(
                "No column titles found on row %s. Check the header row "
                "number.", self.header_row))
        self.write({
            'detected_columns': ', '.join(present),
            'state': 'map',
        })
        self._guess_mapping(present)
        return self._reopen()

    def _guess_mapping(self, headers):
        """Offer a mapping based on the column titles.

        A guess the user can correct beats an empty form: most plants name
        these columns recognisably, and the ones that do not are exactly the
        cases where the user needs to look anyway.
        """
        hints = {
            'col_date': ('date', 'dt', 'day'),
            'col_shift': ('shift',),
            'col_machine': ('machine', 'workcenter', 'work center', 'm/c'),
            'col_product': ('product', 'item', 'sku', 'description'),
            'col_planned': ('target', 'plan', 'planned'),
            'col_actual': ('actual', 'output', 'produced', 'production', 'qty'),
            'col_rejected': ('reject', 'rework', 'scrap', 'defect'),
            'col_run_hours': ('run hour', 'running', 'run hrs', 'machine hour'),
            'col_downtime': ('downtime', 'down time', 'idle', 'breakdown'),
            'col_manpower': ('manpower', 'operator', 'labour', 'labor', 'men'),
            'col_remarks': ('remark', 'note', 'comment', 'reason'),
        }
        lowered = {h.lower(): h for h in headers}
        values = {}
        for field_name, keywords in hints.items():
            if self[field_name]:
                continue
            for header_lower, header in lowered.items():
                if any(keyword in header_lower for keyword in keywords):
                    values[field_name] = header
                    break
        if values:
            self.write(values)

    # ==================================================================
    # Parsing
    # ==================================================================
    def _parse(self):
        """Turn the sheet into (row_number, values dict, error) tuples."""
        self.ensure_one()
        missing = [
            label for field_name, label in (
                ('col_date', _("Date")), ('col_machine', _("Machine")),
                ('col_product', _("Product")), ('col_actual', _("Produced")))
            if not self[field_name]
        ]
        if missing:
            raise UserError(_(
                "These columns must be mapped before importing: %s",
                ', '.join(missing)))

        headers, rows = self._load_rows()
        index = {header: position for position, header in enumerate(headers)}

        def cell(row, column_name):
            if not column_name or column_name not in index:
                return None
            position = index[column_name]
            return row[position] if position < len(row) else None

        machines = self._machine_lookup()
        products = self._product_lookup()
        shifts = {
            (s.code or '').strip().lower(): s
            for s in self.env['fmes.shift'].search(
                [('company_id', '=', self.company_id.id)])
        }
        shifts.update({
            (s.name or '').strip().lower(): s
            for s in self.env['fmes.shift'].search(
                [('company_id', '=', self.company_id.id)])
        })

        parsed = []
        for offset, row in enumerate(rows):
            row_number = self.header_row + 1 + offset
            if row is None or all(value in (None, '') for value in row):
                continue

            raw_date = cell(row, self.col_date)
            date_value = self._coerce_date(raw_date)
            if not date_value:
                parsed.append((row_number, None,
                               _("unreadable date %r") % (raw_date,)))
                continue

            machine_key = self._key(cell(row, self.col_machine))
            machine = machines.get(machine_key)
            if not machine:
                parsed.append((row_number, None,
                               _("unknown machine %r") % (
                                   cell(row, self.col_machine),)))
                continue

            product_key = self._key(cell(row, self.col_product))
            product = products.get(product_key)
            if not product:
                parsed.append((row_number, None,
                               _("unknown product %r") % (
                                   cell(row, self.col_product),)))
                continue

            shift = shifts.get(self._key(cell(row, self.col_shift))) \
                or self.default_shift_id
            if not shift:
                parsed.append((row_number, None, _("no shift on the row and "
                                                  "no default shift set")))
                continue

            actual = self._coerce_float(cell(row, self.col_actual))
            rejected = self._coerce_float(cell(row, self.col_rejected))
            if rejected > actual:
                parsed.append((row_number, None, _(
                    "rejected (%(rejected)s) exceeds produced (%(actual)s)",
                    rejected=rejected, actual=actual)))
                continue

            parsed.append((row_number, {
                'date': date_value,
                'shift_id': shift.id,
                'workcenter_id': machine.id,
                'product_id': product.id,
                'planned_qty': self._coerce_float(cell(row, self.col_planned)),
                'actual_qty': actual,
                'rejected_qty': rejected,
                'run_hours': self._coerce_float(cell(row, self.col_run_hours)),
                'downtime_hours': self._coerce_float(
                    cell(row, self.col_downtime)),
                'actual_manpower': self._coerce_float(
                    cell(row, self.col_manpower)),
                'note': self._coerce_text(cell(row, self.col_remarks)),
                'company_id': self.company_id.id,
            }, None))
        return parsed

    def _machine_lookup(self):
        machines = self.env['mrp.workcenter'].search([])
        lookup = {}
        for machine in machines:
            for value in (machine.fmes_machine_code, machine.code,
                          machine.name):
                key = self._key(value)
                if key and key not in lookup:
                    lookup[key] = machine
        return lookup

    def _product_lookup(self):
        products = self.env['product.product'].search(
            [('type', '!=', 'service')])
        lookup = {}
        for product in products:
            for value in (product.default_code, product.name):
                key = self._key(value)
                if key and key not in lookup:
                    lookup[key] = product
        return lookup

    @staticmethod
    def _key(value):
        if value is None:
            return ''
        return str(value).strip().lower()

    @staticmethod
    def _coerce_float(value):
        if value in (None, ''):
            return 0.0
        try:
            return float(str(value).replace(',', '').strip())
        except (TypeError, ValueError):
            return 0.0

    @staticmethod
    def _coerce_text(value):
        if value in (None, ''):
            return False
        return str(value).strip()

    @staticmethod
    def _coerce_date(value):
        if value in (None, ''):
            return None
        if isinstance(value, str):
            text = value.strip()
            for pattern in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y', '%m/%d/%Y',
                            '%d.%m.%Y', '%d-%b-%Y', '%d %b %Y'):
                try:
                    return fields.Date.to_date(
                        datetime.strptime(text, pattern))
                except ValueError:
                    continue
            return None
        try:
            return fields.Date.to_date(value)
        except (TypeError, ValueError):
            return None

    # ==================================================================
    # Preview and import
    # ==================================================================
    def action_preview(self):
        """Step 2: validate everything and show what would happen."""
        self.ensure_one()
        parsed = self._parse()
        errors = [(number, reason) for number, _vals, reason in parsed
                  if reason]
        valid = [item for item in parsed if item[2] is None]

        rows_html = []
        for number, values, reason in parsed[:200]:
            if reason:
                rows_html.append(
                    '<tr style="background:#fdf2f3">'
                    '<td>%s</td><td colspan="5">%s</td></tr>'
                    % (number, reason))
            else:
                machine = self.env['mrp.workcenter'].browse(
                    values['workcenter_id'])
                product = self.env['product.product'].browse(
                    values['product_id'])
                rows_html.append(
                    '<tr><td>%s</td><td>%s</td><td>%s</td><td>%s</td>'
                    '<td style="text-align:right">%.2f</td>'
                    '<td style="text-align:right">%.2f</td></tr>'
                    % (number, values['date'], machine.display_name,
                       product.display_name, values['planned_qty'],
                       values['actual_qty']))

        html = (
            '<table class="table table-sm">'
            '<thead><tr><th>Row</th><th>Date</th><th>Machine</th>'
            '<th>Product</th><th style="text-align:right">Target</th>'
            '<th style="text-align:right">Produced</th></tr></thead>'
            '<tbody>%s</tbody></table>' % ''.join(rows_html))
        if len(parsed) > 200:
            html += '<p class="text-muted">Showing the first 200 rows.</p>'

        self.write({
            'state': 'preview',
            'row_count': len(parsed),
            'valid_count': len(valid),
            'error_count': len(errors),
            'preview_html': html,
        })
        return self._reopen()

    def action_import(self):
        """Step 3: create the entries as a reversible batch."""
        self.ensure_one()
        parsed = self._parse()
        valid = [(number, values) for number, values, reason in parsed
                 if reason is None]
        if not valid:
            raise UserError(_(
                "Nothing to import: every row was rejected. Check the preview "
                "for the reasons."))

        batch = self.env['fmes.import.batch'].create({
            'source_filename': self.file_name,
            'row_count': len(parsed),
            'skipped_count': len(parsed) - len(valid),
            'company_id': self.company_id.id,
        })

        Entry = self.env['fmes.production.entry']
        created, skipped = 0, []
        for number, values in valid:
            existing = Entry.search([
                ('date', '=', values['date']),
                ('shift_id', '=', values['shift_id']),
                ('workcenter_id', '=', values['workcenter_id']),
                ('product_id', '=', values['product_id']),
                ('company_id', '=', values['company_id']),
            ], limit=1)
            if existing:
                skipped.append(_(
                    "Row %(row)s: an entry already exists for that machine, "
                    "product and shift (%(entry)s).",
                    row=number, entry=existing.name))
                continue
            if not self.create_missing_entries:
                # Only accept rows that match a released plan line.
                probe = Entry.new({
                    'date': values['date'], 'shift_id': values['shift_id'],
                    'workcenter_id': values['workcenter_id'],
                    'product_id': values['product_id'],
                })
                if not probe._find_plan_line():
                    skipped.append(_(
                        "Row %(row)s: no released plan covers that slot.",
                        row=number))
                    continue
            Entry.create(dict(values, import_batch_id=batch.id))
            created += 1

        log_lines = [_(
            "Imported %(created)s of %(total)s rows from %(file)s.",
            created=created, total=len(parsed),
            file=self.file_name or _('the uploaded file'))]
        log_lines += [reason for _number, _vals, reason in parsed if reason]
        log_lines += skipped
        batch.write({
            'log': '\n'.join(log_lines),
            'skipped_count': len(parsed) - created,
        })
        _logger.info("Furnishing MES import %s: %s entries created",
                     batch.name, created)

        return {
            'type': 'ir.actions.act_window',
            'name': _('Import Batch'),
            'res_model': 'fmes.import.batch',
            'res_id': batch.id,
            'view_mode': 'form',
        }

    def action_back_to_mapping(self):
        self.state = 'map'
        return self._reopen()

    def _reopen(self):
        return {
            'type': 'ir.actions.act_window',
            'res_model': self._name,
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }
